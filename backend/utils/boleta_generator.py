"""
Generación de boletas en DOCX (rescate slice 1).
"""
from docx import Document


def generar_boleta(alumno, calificaciones, carrera, config_db):
    """Construye el documento DOCX de la boleta con el historial completo.

    Args:
        alumno: modelo Alumno con `nombre_completo` y `numero_control`.
        calificaciones: lista de Calificacion con `calificacion_final`, `periodo`, `anio` y `materia`.
        carrera: modelo Carrera o None.
        config_db: dict de configuración (claves opcionales `institucion_nombre`, `ciclo_texto`).

    Returns:
        docx.document.Document listo para `doc.save(buffer)`.
    """
    cfg = config_db or {}
    doc = Document()
    doc.add_heading(cfg.get('institucion_nombre', 'Boleta de calificaciones'), level=1)
    doc.add_paragraph(f"Alumno: {alumno.nombre_completo} ({alumno.numero_control or 's/n'})")
    doc.add_paragraph(f"Carrera: {carrera.nombre if carrera else ''}")
    if cfg.get('ciclo_texto'):
        doc.add_paragraph(cfg['ciclo_texto'])
    tabla = doc.add_table(rows=1, cols=4)
    encabezado = tabla.rows[0].cells
    encabezado[0].text, encabezado[1].text, encabezado[2].text, encabezado[3].text = (
        'Materia', 'Calificación', 'Periodo', 'Año')
    for cal in calificaciones:
        fila = tabla.add_row().cells
        fila[0].text = cal.materia.nombre if cal.materia else 'N/A'
        fila[1].text = str(cal.calificacion_final)
        fila[2].text = str(cal.periodo or '')
        fila[3].text = str(cal.anio or '')
    notas = [c.calificacion_final for c in calificaciones if c.calificacion_final and c.calificacion_final > 0]
    promedio = round(sum(notas) / len(notas), 1) if notas else 0
    doc.add_paragraph(f'Promedio: {promedio} ({len(notas)} materias)')
    return doc
